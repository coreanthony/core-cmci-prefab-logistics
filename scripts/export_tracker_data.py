"""Export the logistics workbook into the static tracker data contract."""

from __future__ import annotations

import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "deliverables" / "Pre-fab Consultation Room - Core CMCI Logistics Plan.xlsx"
OUTPUT = ROOT / "site-data" / "tracker.json"


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def business_days_before(value: date, count: int) -> date:
    cursor = value
    remaining = int(count)
    while remaining:
        cursor -= timedelta(days=1)
        if cursor.weekday() < 5:
            remaining -= 1
    return cursor


def clean(value, fallback=""):
    if value is None:
        return fallback
    return str(value).strip()


def main():
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)

    assumptions_ws = wb["Assumptions"]
    default_units = int(assumptions_ws["B4"].value)
    delivery_buffer = int(assumptions_ws["B5"].value)
    install_days = int(assumptions_ws["B6"].value)
    transit_days = {
        clean(assumptions_ws.cell(row=row, column=1).value): int(assumptions_ws.cell(row=row, column=2).value)
        for row in range(11, 17)
    }

    map_ws = wb["Map Data"]
    map_headers = [cell.value for cell in next(map_ws.iter_rows(min_row=1, max_row=1))]
    map_by_cs = {}
    for row in map_ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(map_headers, row))
        map_by_cs[clean(record["CS#"])] = record

    current_ws = wb["Current List"]
    headers = [current_ws.cell(row=1, column=column).value for column in range(1, 23)]
    stores = []
    for row in current_ws.iter_rows(min_row=2, max_col=22, values_only=True):
        source = dict(zip(headers, row))
        cs_number = clean(source["CS#"])
        mapped = map_by_cs[cs_number]
        region = clean(mapped["Region"])
        msd = source["MSD"].date()
        csd = source["CSD"]
        if not isinstance(csd, (datetime, date)):
            csd = msd + timedelta(days=7)
        elif isinstance(csd, datetime):
            csd = csd.date()

        raw_config = source["Unit Configuration"]
        if isinstance(raw_config, (int, float)) and raw_config > 0:
            parsed_units = int(raw_config)
        else:
            match = re.match(r"^\s*(\d+)", clean(raw_config))
            parsed_units = int(match.group(1)) if match else None
        config_confirmed = parsed_units is not None and parsed_units > 0
        units = parsed_units if config_confirmed else default_units
        layout_status = clean(source["Layout Validation Status"], "TBD")
        labor = clean(source["MV or Store Labor?"], "TBD")
        if labor in {"0", "None", ""}:
            labor = "TBD"

        deliver_by = business_days_before(msd, delivery_buffer)
        ship_by = business_days_before(deliver_by, transit_days[region])
        ship_week = ship_by - timedelta(days=ship_by.weekday())

        flags = []
        if layout_status != "Approved":
            flags.append("Layout not validated")
        if not config_confirmed:
            flags.append("Config TBD; default units used")
        if labor == "TBD":
            flags.append("Labor TBD")
        if region == "Outlier":
            flags.append("Remote / out of route")
        if region == "South TX":
            flags.append("Long haul")

        location_basis = clean(mapped["Location basis"])
        if "approx" in location_basis.lower() or "verify" in location_basis.lower():
            flags.append("Location verification required")

        stores.append(
            {
                "region": region,
                "cluster": clean(source["Cluster"]),
                "csNumber": cs_number,
                "storeNumber": clean(source["5 Digit Store"] or source["Store #"]).zfill(5),
                "address": clean(source["Address"]),
                "city": clean(source["City"]).title(),
                "state": clean(source["State"]),
                "zip": clean(source["Zip"]).zfill(5),
                "storeType": clean(source["Store Type"]),
                "projectManager": clean(source["PM/SPM"]),
                "layoutStatus": layout_status,
                "laborModel": labor,
                "merchandisingHours": source["Merchandising Hours"] or 0,
                "mvVendor": clean(source["MV"], "-") if clean(source["MV"]) != "N/A" else "-",
                "configRaw": raw_config if config_confirmed else None,
                "units": units,
                "configStatus": "Confirmed" if config_confirmed else "TBD",
                "msd": msd.isoformat(),
                "deliverBy": deliver_by.isoformat(),
                "shipBy": ship_by.isoformat(),
                "shipWeek": ship_week.isoformat(),
                "csd": csd.isoformat(),
                "latitude": mapped["Latitude"],
                "longitude": mapped["Longitude"],
                "locationBasis": location_basis,
                "flags": flags,
            }
        )

    stores.sort(key=lambda item: (item["msd"], item["region"], item["city"], item["storeNumber"]))

    open_items_ws = wb["Open Items"]
    open_items = []
    for item, detail, owner in open_items_ws.iter_rows(min_row=4, max_col=3, values_only=True):
        if not item:
            continue
        open_items.append(
            {
                "item": clean(item),
                "detail": clean(detail),
                "owner": clean(owner).replace("Owner:", "").strip(),
            }
        )

    payload = {
        "meta": {
            "program": "CVS Pre-fab Consultation Room",
            "contractor": "Core CMCI",
            "sourceWorkbook": WORKBOOK.name,
            "sourceDate": "2026-09-25",
            "exportedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "recordCount": len(stores),
        },
        "assumptions": {
            "defaultUnits": default_units,
            "deliveryBufferBusinessDays": delivery_buffer,
            "installDaysPerStore": install_days,
            "transitBusinessDays": transit_days,
            "status": "Placeholder assumptions; vendor ship-from point and lead time remain open.",
        },
        "openItems": open_items,
        "stores": stores,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Exported {len(stores)} stores to {OUTPUT}")


if __name__ == "__main__":
    main()
