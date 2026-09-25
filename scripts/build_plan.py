import openpyxl, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
src='/mnt/user-data/outputs/Pre-fab Consultation Room - Core CMCI Only.xlsx'
d=pd.read_excel(src,sheet_name=0,dtype={'CS#':str})
reg={'DFW':['CARROLLTON','COPPELL','CORINTH','DALLAS','DENTON','FLOWER MOUND','FRISCO','GRAND PRAIRIE','IRVING','LEWISVILLE','MCKINNEY','PLANO','ROYSE CITY'],
'Houston':['HOUSTON','KATY','CYPRESS','KINGWOOD','LEAGUE CITY','MANVEL','PASADENA','PEARLAND','MISSOURI CITY','ROSENBERG','RICHMOND','SUGAR LAND','TOMBALL','WILLIS','MONTGOMERY','HUNTSVILLE'],
'Austin / Central':['AUSTIN','ROUND ROCK','SAN MARCOS','MARBLE FALLS','KILLEEN'],
'San Antonio':['SAN ANTONIO','ALAMO HEIGHTS','CIBOLO','NEW BRAUNFELS','SEGUIN'],
'South TX':['EDINBURG','MISSION','WESLACO','CORPUS CHRISTI'],
'Outlier':['AMARILLO','NEW ORLEANS']}
m={c:r for r,cs in reg.items() for c in cs}
d['Region']=d['City'].map(m); assert d['Region'].notna().all()
order=list(reg); d['ro']=d['Region'].map(order.index)
d=d.sort_values(['ro','MSD','City','Store #']).reset_index(drop=True)

wb=openpyxl.load_workbook(src)
for n in ['Assumptions','Logistics Plan','Region Summary','Open Items']:
    if n in wb.sheetnames: del wb[n]
hf=Font(bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F3864')
inp=Font(color='0000FF'); infill=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BFBFBF'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,row,c1,c2):
    for c in range(c1,c2+1):
        x=ws.cell(row,c); x.font=hf; x.fill=hfill; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); x.border=bd

# Assumptions
A=wb.create_sheet('Assumptions')
A['A1']='Logistics Plan - Assumptions (edit yellow cells; plan recalculates)'; A['A1'].font=Font(bold=True,size=13)
A['A3']='Global inputs'; A['A3'].font=Font(bold=True)
rows=[('Default units per store when config is blank/TBD',3,'Most common configuration among the 35 validated stores (29 are 3-unit)'),
('Delivery buffer before MSD (business days)',3,'Units on site, checked and staged before install starts. ASSUMPTION - confirm with vendor/CVS'),
('Install days per store (crew-days)',1,'ASSUMPTION - confirm actual install duration per unit config')]
for i,(a,b,c) in enumerate(rows,4):
    A.cell(i,1,a); x=A.cell(i,2,b); x.font=inp; x.fill=infill; A.cell(i,3,c)
A['A9']='Transit time from vendor/staging point to region (business days)'; A['A9'].font=Font(bold=True)
for j,h in enumerate(['Region','Transit days','Note'],1): A.cell(10,j,h)
hdr(A,10,1,3)
tr={'DFW':(2,'PLACEHOLDER - update once ship-from point is known'),'Houston':(2,'PLACEHOLDER'),'Austin / Central':(2,'PLACEHOLDER'),'San Antonio':(2,'PLACEHOLDER'),'South TX':(3,'PLACEHOLDER - RGV / Corpus are long hauls'),'Outlier':(4,'PLACEHOLDER - Amarillo and New Orleans are out of the way')}
for i,(r,(t,n)) in enumerate(tr.items(),11):
    A.cell(i,1,r); x=A.cell(i,2,t); x.font=inp; x.fill=infill; A.cell(i,3,n)
A['A18']='Definitions used'; A['A18'].font=Font(bold=True)
for i,t in enumerate(['MSD = date install window opens per the store list. CSD = MSD + 7 days (existing formula in the list).',
 'Deliver-by = MSD less the delivery buffer (business days). Ship-by = Deliver-by less regional transit days.',
 'Units come from the Unit Configuration column; stores not yet validated have none, so the default above is used and flagged TBD.',
 'Regions are grouped by metro so one crew/truck run covers a cluster. Killeen is grouped with Austin / Central.',
 'Plan tab pulls MSD, labor and status live from Current List by CS#.'],19):
    A.cell(i,1,t)
A.column_dimensions['A'].width=58; A.column_dimensions['B'].width=14; A.column_dimensions['C'].width=90

# Logistics Plan
P=wb.create_sheet('Logistics Plan')
P['A1']='Core CMCI - Consultation Room Logistics Plan'; P['A1'].font=Font(bold=True,size=13)
P['A2']='Sorted by region, then MSD. Dates and units are formulas driven by the Assumptions tab and Current List.'
heads=['Region','CS#','Store #','City','State','Zip','Store Type','Labor (MV / Store)','MV Vendor','Layout Status','Config (raw)','Units','Config','MSD','Deliver-by','Ship-by','Ship week of (Mon)','CSD','Flags']
HR=4
for j,h in enumerate(heads,1): P.cell(HR,j,h)
hdr(P,HR,1,len(heads)); P.row_dimensions[HR].height=32
CL="'Current List'!"
def lk(col,r): return f"INDEX({CL}${col}:${col},MATCH($B{r},{CL}$A:$A,0))"
for i,row in d.iterrows():
    r=HR+1+i
    P.cell(r,1,row['Region']); P.cell(r,2,row['CS#']); P.cell(r,3,int(row['Store #'])); P.cell(r,4,row['City'].title()); P.cell(r,5,row['State']); P.cell(r,6,str(row['Zip']))
    P.cell(r,7,f'={lk("L",r)}')
    P.cell(r,8,f'=IF({lk("O",r)}=0,"TBD",{lk("O",r)})')
    P.cell(r,9,f'=IF(OR({lk("R",r)}=0,{lk("R",r)}="N/A"),"-",{lk("R",r)})')
    P.cell(r,10,f'={lk("N",r)}')
    P.cell(r,11,f'=IF({lk("S",r)}=0,"",{lk("S",r)})')
    P.cell(r,12,f'=IFERROR(IF(ISNUMBER(K{r}),K{r},VALUE(LEFT(K{r},1))),Assumptions!$B$4)')
    P.cell(r,13,f'=IF(K{r}="","TBD","Confirmed")')
    P.cell(r,14,f'={lk("T",r)}'); 
    P.cell(r,15,f'=WORKDAY(N{r},-Assumptions!$B$5)')
    P.cell(r,16,f'=WORKDAY(O{r},-VLOOKUP(A{r},Assumptions!$A$11:$B$16,2,FALSE))')
    P.cell(r,17,f'=P{r}-WEEKDAY(P{r},3)')
    P.cell(r,18,f'={lk("U",r)}')
    P.cell(r,19,f'=TRIM(IF(J{r}<>"Approved","Layout not validated; ","")&IF(M{r}="TBD","Config TBD (default units); ","")&IF(H{r}="TBD","Labor TBD; ","")&IF(A{r}="Outlier","Remote / out of route; ","")&IF(A{r}="South TX","Long haul; ",""))')
    for c in range(1,20): P.cell(r,c).border=bd
    for c in (14,15,16,17,18): P.cell(r,c).number_format='ddd m/d/yy'
last=HR+len(d)
widths=[16,9,8,16,6,7,20,15,16,17,15,7,11,13,13,13,15,13,44]
for j,w in enumerate(widths,1): P.column_dimensions[L(j)].width=w
P.freeze_panes=P.cell(HR+1,5); P.auto_filter.ref=f'A{HR}:{L(len(heads))}{last}'
tr_=last+1
P.cell(tr_,1,'Total').font=Font(bold=True); P.cell(tr_,3,f'=COUNTA(C{HR+1}:C{last})').font=Font(bold=True); P.cell(tr_,12,f'=SUM(L{HR+1}:L{last})').font=Font(bold=True)
yel=PatternFill('solid',fgColor='FCE4D6')
from openpyxl.formatting.rule import FormulaRule
P.conditional_formatting.add(f'S{HR+1}:S{last}',FormulaRule(formula=[f'LEN(S{HR+1})>0'],fill=yel))

# Region Summary
S=wb.create_sheet('Region Summary')
S['A1']='Region x MSD week: stores and units'; S['A1'].font=Font(bold=True,size=13)
weeks=sorted(d['MSD'].unique())
S['A3']='Stores'; S['A3'].font=Font(bold=True)
def block(top,kind):
    S.cell(top,1,'Region')
    for j,w in enumerate(weeks,2):
        S.cell(top,j,pd.Timestamp(w).to_pydatetime()).number_format='m/d/yy'
    S.cell(top,len(weeks)+2,'Total')
    hdr(S,top,1,len(weeks)+2)
    for i,rg in enumerate(order,top+1):
        S.cell(i,1,rg)
        for j in range(2,len(weeks)+2):
            c=L(j)
            f=(f"=COUNTIFS('Logistics Plan'!$A${HR+1}:$A${last},$A{i},'Logistics Plan'!$N${HR+1}:$N${last},{c}${top})" if kind=='n' else
               f"=SUMIFS('Logistics Plan'!$L${HR+1}:$L${last},'Logistics Plan'!$A${HR+1}:$A${last},$A{i},'Logistics Plan'!$N${HR+1}:$N${last},{c}${top})")
            S.cell(i,j,f)
        S.cell(i,len(weeks)+2,f'=SUM(B{i}:{L(len(weeks)+1)}{i})')
    t=top+len(order)+1
    S.cell(t,1,'Total').font=Font(bold=True)
    for j in range(2,len(weeks)+3): S.cell(t,j,f'=SUM({L(j)}{top+1}:{L(j)}{t-1})').font=Font(bold=True)
    for r in range(top+1,t+1):
        for c in range(1,len(weeks)+3): S.cell(r,c).border=bd
    return t
t1=block(4,'n')
S.cell(t1+2,1,'Units (default used where config TBD)').font=Font(bold=True)
t2=block(t1+3,'u')
cr=t2+2
S.cell(cr,1,'Crew-days needed (stores x install days per store)').font=Font(bold=True)
S.cell(cr+1,1,'Peak week crew-days'); S.cell(cr+1,2,f'=MAX(B{t1}:{L(len(weeks)+1)}{t1})*Assumptions!$B$6')
S.cell(cr+2,1,'Total crew-days'); S.cell(cr+2,2,f'={L(len(weeks)+2)}{t1}*Assumptions!$B$6')
S.cell(cr+3,1,'Crews needed at peak if each does 5 install days/week'); S.cell(cr+3,2,f'=ROUNDUP(B{cr+1}/5,0)')
S.column_dimensions['A'].width=48
for j in range(2,len(weeks)+3): S.column_dimensions[L(j)].width=11

# Open Items
O=wb.create_sheet('Open Items')
O['A1']='Open items to close before the plan is final'; O['A1'].font=Font(bold=True,size=13)
items=[('Ship-from point / vendor lead time','Transit days on Assumptions are placeholders. Need the vendor ship point and production lead time.','Owner: Procurement / vendor'),
('Meaning and firmness of MSD','Plan treats MSD as the install window opening and CSD (MSD + 7) as the deadline. Confirm with CVS/program manager.','Owner: PM'),
('Unit configs for stores not validated','30 stores show Not Validated Yet and have no unit config. Plan uses the default and flags them; PO quantities should not be locked until configs are confirmed.','Owner: CVS / layout validation'),
('Labor model TBD','MV vs Store Labor is blank for the not-validated stores and some validated stores (hours 0). Affects who receives/installs.','Owner: PM'),
('Remote stores','Amarillo (1543) and New Orleans (10594) are far from any cluster; decide on dedicated delivery or bundling with another run.','Owner: Logistics'),
('Receiving windows / store contacts','Need site delivery hours, dock/no-dock, and receiving contacts per store, particularly for freestanding pads and convenience store types.','Owner: Field ops'),
('Crew plan','Region Summary gives peak-week load; confirm self-perform vs sub and travel/lodging for South TX and outliers.','Owner: COO / Ops'),
('Tie to PO','Once configs firm up, roll units per region into the vendor PO and delivery schedule.','Owner: Procurement')]
for j,h in enumerate(['Item','Detail','Owner'],1): O.cell(3,j,h)
hdr(O,3,1,3)
for i,it in enumerate(items,4):
    for j,v in enumerate(it,1):
        c=O.cell(i,j,v); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=bd
O.column_dimensions['A'].width=38; O.column_dimensions['B'].width=95; O.column_dimensions['C'].width=32
wb.move_sheet('Rejected Candidates',offset=0)
out='/mnt/user-data/outputs/Pre-fab Consultation Room - Core CMCI Logistics Plan.xlsx'
wb.save(out); print(out)
